# -*- coding: utf-8 -*-
"""Excel รายเทรนด์ x หมวดสินค้า (export_trend_categories.py, 🆕 2026-10-29) - ข้อมูลสังเคราะห์เล็กๆ ไม่เรียก API

รัน: pytest tool/tests/test_trend_categories_export.py -v
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "phase3_supply_layer"))
for _p in Path(__file__).resolve().parents:
    if (_p / "common" / "bootstrap.py").exists():
        sys.path.insert(0, str(_p))
        break

import pandas as pd

import export_trend_categories as ex


def _ctx():
    return {"meta": {"run_id": "r1", "topic": "beauty", "n_trends_total": 3, "data_as_of": "Sep 2026"},
            "trends": [{
                "rank": 1, "name": "Sun Trend", "key": "Search::Sun Trend",
                "growth": {"scores": [0.9, 0.8, 0.88, 0.7, 0.95], "past_growth_pct": 114.9,
                           "forecast_pct": [14.7, 40.3, 79.4, 98.9, 127.1], "forecast_cagr": [14.7, 18.45, 21.509, 18.7, 17.827]},
                "category_agreement": {"category": "Sun Care", "in_main": 3},
                "beauty_need": {"key_benefits": "sun protection", "product_keywords": ["face sunscreen", "tinted"],
                                "product_categories": ["Skin Care", "Sun Care"]},
                "l2": {"category": "Sun Care", "n": 1062}, "no_close_product": False, "similarity_stats": {"max": 0.355},
                "talk_sources": {"categories": ["Health"]}, "places": ["TH", "WW"], "keywords": ["best sunscreen"],
                "products": [{"name_full": "KA Sunscreen 50g", "product_url": "https://shop/ka"}]}]}


def _result():
    return {"n_products_total": 22392, "catalog_filter": "L1 = Beauty and Personal Care",
            "results": {"Search::Sun Trend": {
                "main_category": {"category": "Sun Care", "reason": "sun focus"},
                "top_matches": [{"product_name": "KA Sunscreen 50g", "brand": "KA", "category": "Sun Care", "subgroup": "Face Sunscreen",
                                 "similarity": 0.355, "sale_price_thb": 229.0, "unit_price": 4.58333, "unit_label": "g"},
                                {"product_name": "Toothpaste 120g", "brand": "Vussen", "category": "Oral Care", "subgroup": "Toothpaste",
                                 "similarity": 0.34, "sale_price_thb": 249.0, "unit_price": 2.07, "unit_label": "g"},
                                {"product_name": "Sun Body 100ml", "brand": "B", "category": "Sun Care", "subgroup": "Body Sunscreen",
                                 "similarity": 0.33, "sale_price_thb": 199.0, "unit_price": 1.99, "unit_label": "ml"}],
                "by_category": {"pool_size": 224, "main_category": "Sun Care", "n_category": 1062,
                                "distribution": [{"group": "Face Sunscreen", "n": 200}, {"group": "Body Sunscreen", "n": 20},
                                                 {"group": None, "n": 4}],
                                "categories": [{"subgroup": "Face Sunscreen", "n_total": 881, "median_price_thb": 289.0,
                                                "min_price_thb": 15.0, "max_price_thb": 2950.0,
                                                "products": [{"product_name": "KA Sunscreen 50g", "brand": "KA",
                                                              "similarity": 0.355, "rank_in_market": 1}]}]}}}}


def test_helpers_summarise_types_and_nearest_categories():
    by_cat = _result()["results"]["Search::Sun Trend"]["by_category"]
    assert ex.top_types(by_cat) == "Face Sunscreen 200; Body Sunscreen 20"           # ก้อน "อื่นๆ" (None) ไม่นับ
    assert ex.nearest_categories(_result()["results"]["Search::Sun Trend"]["top_matches"]) == "Sun Care 2; Oral Care 1"
    assert ex.top_types(None) == "" and ex.nearest_categories(None) == ""


def test_summary_row_matches_report_numbers_and_rounds():
    row = ex.summary_rows(_ctx(), _result())[0]
    assert row["คะแนนปี 3 (0-1)"] == 0.88 and row["คะแนนปี 5 (0-1)"] == 0.95
    assert row["คาดโตสะสมปี 3 (%)"] == 79.4 and row["คาดโตเฉลี่ยต่อปี ถึงปี 3 (%)"] == 21.5     # ปัดเหมือนตัวเลขในรายงาน
    assert row["หมวดหลัก (L2)"] == "Sun Care" and row["สินค้าในหมวดหลักทั้งร้าน (ชิ้น)"] == 1062 and row["กลุ่มเทียบ (ชิ้น)"] == 224
    assert row["สินค้าที่ใกล้พอ"] == "มี" and row["หมวดสินค้าที่ AI โยงไว้"] == "Skin Care; Sun Care"
    ctx = _ctx()
    ctx["trends"][0]["no_close_product"] = True
    assert ex.summary_rows(ctx, _result())[0]["สินค้าที่ใกล้พอ"] == "ไม่มีสินค้าที่ใกล้พอ"


def test_type_rows_share_and_report_flag():
    rows = ex.type_rows(_ctx(), _result())
    assert [r["ชนิดสินค้า (L4)"] for r in rows] == ["Face Sunscreen", "Body Sunscreen", ex.OTHER_TYPES]
    face, body, other = rows
    assert face["สัดส่วนในกลุ่มเทียบ (%)"] == 89.3 and face["แสดงในรายงาน"] == "ใช่" and face["ราคากลาง (฿)"] == 289.0
    assert face["สินค้าที่ใกล้เทรนด์สุดในชนิดนี้"] == "KA Sunscreen 50g" and face["อันดับในตลาด"] == 1
    assert body["แสดงในรายงาน"] == "ไม่" and body["สินค้าทั้งร้านในชนิดนี้ (ชิ้น)"] is None     # ชนิดที่ไม่แสดงในรายงานไม่มีราคา/สินค้า


def test_product_rows_links_come_from_report_then_catalog():
    rows = ex.product_rows(_ctx(), _result(), urls={("Vussen", "Toothpaste 120g"): "https://shop/tp"})
    assert [r["ลำดับสินค้า"] for r in rows] == [1, 2, 3]
    assert rows[0]["ลิงก์สินค้า"] == "https://shop/ka" and rows[1]["ลิงก์สินค้า"] == "https://shop/tp"
    assert rows[2]["ลิงก์สินค้า"] is None and rows[0]["฿ ต่อหน่วย"] == 4.58


def test_google_rows_split_category_and_topic():
    cands = pd.DataFrame([{"trend_name": "Sun Trend",
                           "source_themes": "Health: Sunscreen and Skin Care | Beauty & Fitness: Skincare: Basics"}])
    rows = ex.google_rows(_ctx(), cands)
    assert [(r["หมวดของ Google"], r["หัวข้อ"]) for r in rows] == [
        ("Health", "Sunscreen and Skin Care"), ("Beauty & Fitness", "Skincare: Basics")]   # ตัดที่ ": " แรกเท่านั้น
    assert ex.google_rows(_ctx(), None) == []


def test_write_excel_roundtrip_with_links(tmp_path):
    frame = pd.DataFrame([{"สินค้า": "A", "ลิงก์สินค้า": "https://shop/a"}, {"สินค้า": "B", "ลิงก์สินค้า": None}])
    path = ex.write_excel(tmp_path / "x.xlsx", [("ชีต", frame)])
    back = pd.read_excel(path, sheet_name="ชีต")
    assert list(back["สินค้า"]) == ["A", "B"] and back["ลิงก์สินค้า"].iloc[0] == "https://shop/a"
    from openpyxl import load_workbook
    ws = load_workbook(path)["ชีต"]
    assert ws.freeze_panes == "A2" and ws["B2"].hyperlink is not None
