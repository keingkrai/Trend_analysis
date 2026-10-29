# -*- coding: utf-8 -*-
"""
⑫+ Excel รายเทรนด์ x หมวดสินค้า (🆕 2026-10-29, เจ้าของงานสั่ง) - เปิดดูในหนึ่งไฟล์ว่าแต่ละเทรนด์ในรายงาน ⑯ อยู่ในหมวดสินค้าอะไร

อ่านผลที่มีอยู่แล้วของรอบ ไม่เรียก API ไม่คำนวณใหม่ (ตัวเลขเดียวกับรายงาน ⑯):
  phase6_report_context.json            เทรนด์ในรายงาน + คะแนน/% ที่คาด + ความต้องการความงามที่ AI อนุมาน + ธง "ไม่มีสินค้าที่ใกล้พอ"
  phase3_scale_validation_nvidia_result.json   (⑫) หมวดหลัก L2 + เหตุผล + ชนิดสินค้าในกลุ่มเทียบ + สินค้า 15 ชิ้นที่ใกล้ที่สุด
  phase2_search_candidates.csv          (ไม่บังคับ) หัวข้อต้นทางของแต่ละเทรนด์ตามหมวดของ Google
เขียน <รอบ>/phase3_trend_categories.xlsx 5 ชีต: สรุปรายเทรนด์ / ชนิดสินค้าในหมวดหลัก / สินค้าใกล้เทรนด์สุด / หมวด Google ต้นทาง / คำอธิบาย

รัน (จากโฟลเดอร์บนสุดของ repo):
  python tool/phase3_supply_layer/export_trend_categories.py --run <ชื่อรอบ>     (ไม่ใส่ = TREND_RUN_ID หรือรอบล่าสุดที่มี context)
"""
import argparse
import json
import os
import sys
from collections import Counter
from pathlib import Path

import pandas as pd

for _p in Path(__file__).resolve().parents:
    if (_p / "common" / "bootstrap.py").exists():
        sys.path.insert(0, str(_p))
        break

CONTEXT_FILE = "phase6_report_context.json"
RESULT_FILE = "phase3_scale_validation_nvidia_result.json"
CANDIDATES_FILE = "phase2_search_candidates.csv"
OUTPUT_FILE = "phase3_trend_categories.xlsx"

SHEET_SUMMARY, SHEET_TYPES, SHEET_PRODUCTS = "สรุปรายเทรนด์", "ชนิดสินค้าในหมวดหลัก", "สินค้าใกล้เทรนด์สุด"
SHEET_GOOGLE, SHEET_NOTES = "หมวด Google ต้นทาง", "คำอธิบาย"
OTHER_TYPES = "ชนิดอื่นๆ (รวม)"


def _at(values, i):
    """ค่าที่ตำแหน่ง i ของ list (ปี 1..5 = 0..4) - None ถ้าไม่มี"""
    return values[i] if isinstance(values, list) and len(values) > i else None


def _r(v, n=1):
    """ปัดทศนิยม (ตัวเลขในรายงานแสดง 1 ตำแหน่ง) - None คงเดิม"""
    return round(v, n) if isinstance(v, (int, float)) else v


def _join(values, sep="; "):
    return sep.join(str(v) for v in (values or []) if v not in (None, ""))


def _result_of(result, trend):
    return (result.get("results") or {}).get(trend.get("key")) or {}


def top_types(by_category, n=3):
    """ชนิดสินค้าที่เด่นที่สุดในกลุ่มเทียบ เช่น "Face Sunscreen 212; Body Sunscreen 12" (ไม่รวมก้อน "อื่นๆ")"""
    items = [d for d in (by_category or {}).get("distribution", []) if d.get("group")]
    return "; ".join(f"{d['group']} {d['n']}" for d in items[:n])


def nearest_categories(top_matches):
    """หมวด L2 ของสินค้า 15 ชิ้นที่ใกล้เทรนด์ที่สุด เช่น "Sun Care 12; Deodorants 3" """
    counts = Counter(m.get("category") for m in top_matches or [] if m.get("category"))
    return "; ".join(f"{c} {n}" for c, n in counts.most_common())


def summary_rows(ctx, result):
    rows = []
    for t in ctx["trends"]:
        res, g = _result_of(result, t), t.get("growth") or {}
        main = (res.get("main_category") or {})
        agree, need, l2 = t.get("category_agreement") or {}, t.get("beauty_need") or {}, t.get("l2") or {}
        by_cat = res.get("by_category") or {}
        sim = t.get("similarity_stats") or res.get("similarity_stats") or {}
        rows.append({
            "อันดับ": t.get("rank"), "เทรนด์": t.get("name"),
            "คะแนนปี 3 (0-1)": _at(g.get("scores"), 2), "คะแนนปี 5 (0-1)": _at(g.get("scores"), 4),
            "โต 12 เดือนที่ผ่านมา (%)": g.get("past_growth_pct"),
            "คาดโตสะสมปี 1 (%)": _at(g.get("forecast_pct"), 0), "คาดโตสะสมปี 3 (%)": _at(g.get("forecast_pct"), 2),
            "คาดโตสะสมปี 5 (%)": _at(g.get("forecast_pct"), 4),
            "คาดโตเฉลี่ยต่อปี ถึงปี 3 (%)": _r(_at(g.get("forecast_cagr"), 2)),
            "คาดโตเฉลี่ยต่อปี ถึงปี 5 (%)": _r(_at(g.get("forecast_cagr"), 4)),
            "หมวดหลัก (L2)": main.get("category") or agree.get("category") or l2.get("category"),
            "เหตุผลที่เลือกหมวดหลัก": main.get("reason"),
            "สินค้าในหมวดหลักทั้งร้าน (ชิ้น)": l2.get("n") or by_cat.get("n_category"),
            "กลุ่มเทียบ (ชิ้น)": by_cat.get("pool_size"),
            "ชนิดสินค้าที่เด่นในกลุ่มเทียบ": top_types(by_cat),
            "หมวดของสินค้า 15 ชิ้นที่ใกล้เทรนด์สุด": nearest_categories(res.get("top_matches")),
            "Top 5 ที่อยู่ในหมวดหลัก (ชิ้น)": agree.get("in_main"),
            "สินค้าที่ใกล้พอ": "ไม่มีสินค้าที่ใกล้พอ" if t.get("no_close_product") else "มี",
            "ความใกล้เคียงสูงสุด": sim.get("max"),
            "ความต้องการความงามที่ AI อนุมาน": need.get("key_benefits"),
            "คำสินค้าที่ AI อนุมาน": _join(need.get("product_keywords")),
            "หมวดสินค้าที่ AI โยงไว้": _join(need.get("product_categories")),
            "หมวดของ Google ต้นทาง": _join((t.get("talk_sources") or {}).get("categories")),
            "ค้นเพิ่มขึ้นที่": _join(t.get("places"), ", "),
            "คำค้นหลัก": _join(t.get("keywords")),
        })
    return rows


def type_rows(ctx, result):
    """ชนิดสินค้า (Level 4) ของหมวดหลักในกลุ่มเทียบ - ทุกชนิดที่มีใน distribution (ชนิดที่แสดงในรายงานมีราคา/สินค้าที่ใกล้สุดด้วย)"""
    rows = []
    for t in ctx["trends"]:
        by_cat = _result_of(result, t).get("by_category") or {}
        pool = by_cat.get("pool_size") or 0
        shown = {c.get("subgroup"): c for c in by_cat.get("categories", [])}
        for d in by_cat.get("distribution", []):
            c = shown.get(d.get("group")) or {}
            top = (c.get("products") or [{}])[0]
            rows.append({
                "อันดับเทรนด์": t.get("rank"), "เทรนด์": t.get("name"), "หมวดหลัก (L2)": by_cat.get("main_category"),
                "ชนิดสินค้า (L4)": d.get("group") or OTHER_TYPES,
                "ในกลุ่มเทียบ (ชิ้น)": d.get("n"),
                "สัดส่วนในกลุ่มเทียบ (%)": round(100 * d["n"] / pool, 1) if pool else None,
                "แสดงในรายงาน": "ใช่" if c else "ไม่",
                "สินค้าทั้งร้านในชนิดนี้ (ชิ้น)": c.get("n_total"),
                "ราคากลาง (฿)": c.get("median_price_thb"), "ราคาต่ำสุด (฿)": c.get("min_price_thb"),
                "ราคาสูงสุด (฿)": c.get("max_price_thb"),
                "สินค้าที่ใกล้เทรนด์สุดในชนิดนี้": top.get("product_name") or top.get("name"),
                "แบรนด์": top.get("brand"), "คะแนนความใกล้เคียง": top.get("similarity"),
                "อันดับในตลาด": top.get("rank_in_market"),
            })
    return rows


def product_rows(ctx, result, urls=None):
    urls = urls or {}
    rows = []
    for t in ctx["trends"]:
        own = {p.get("name_full") or p.get("name"): p.get("product_url") for p in t.get("products", [])}
        for i, m in enumerate(_result_of(result, t).get("top_matches") or [], start=1):
            name, brand = m.get("product_name"), m.get("brand")
            rows.append({
                "อันดับเทรนด์": t.get("rank"), "เทรนด์": t.get("name"), "ลำดับสินค้า": i,
                "สินค้า": name, "แบรนด์": brand, "หมวด (L2)": m.get("category"), "ชนิดสินค้า (L4)": m.get("subgroup"),
                "คะแนนความใกล้เคียง": m.get("similarity"), "ราคา (฿)": m.get("sale_price_thb"), "ขนาด": m.get("size"),
                "฿ ต่อหน่วย": _r(m.get("unit_price"), 2), "หน่วย": m.get("unit_label"),
                "ขนาดทดลอง/ซอง": "ใช่" if m.get("is_mini") else "", "ขายไปแล้ว": m.get("sold_count"),
                "เหตุผล": m.get("match_reason"),
                "ลิงก์สินค้า": own.get(name) or urls.get((brand, name)),
            })
    return rows


def google_rows(ctx, candidates):
    """หัวข้อต้นทางของเทรนด์ตามหมวดของ Google: source_themes = "หมวด: หัวข้อ | หมวด: หัวข้อ ..." """
    if candidates is None or "source_themes" not in candidates:
        return []
    by_name = {r["trend_name"]: r for r in candidates.to_dict("records")}
    rows = []
    for t in ctx["trends"]:
        for part in str(by_name.get(t.get("name"), {}).get("source_themes") or "").split(" | "):
            if not part.strip():
                continue
            category, _, topic = part.partition(": ")
            rows.append({"อันดับเทรนด์": t.get("rank"), "เทรนด์": t.get("name"),
                         "หมวดของ Google": category.strip(), "หัวข้อ": topic.strip()})
    return rows


def notes_rows(ctx, result):
    meta = ctx.get("meta", {})
    return [
        ("รอบ", meta.get("run_id")), ("หัวข้อ", meta.get("topic")),
        ("เทรนด์ในไฟล์", f"{len(ctx['trends'])} เทรนด์ที่มีหน้ารายละเอียดในรายงาน ⑯ (จากทั้งหมด {meta.get('n_trends_total')})"),
        ("ข้อมูลการค้นหาถึง", meta.get("data_as_of")),
        ("ตลาดที่เทียบ", f"{result.get('n_products_total'):,} สินค้า ({result.get('catalog_filter')}) - หมวด = Level 2 ของ Euromonitor, "
                         "ชนิดสินค้า = Level 4"),
        ("หมวดหลัก (L2)", "AI เลือกหมวด L2 หมวดเดียวต่อเทรนด์ (พร้อมเหตุผล) - รายงานเทียบสินค้าเฉพาะในหมวดนี้"),
        ("กลุ่มเทียบ", "สินค้าที่ใกล้เทรนด์ที่สุดในหมวดหลัก จำนวนประมาณ 1% ของสินค้าทั้งร้าน (คอลัมน์ 'กลุ่มเทียบ (ชิ้น)')"),
        ("ความใกล้เคียง", "ความคล้ายเชิงความหมายระหว่างข้อความของเทรนด์กับชื่อสินค้า (cosine) - ใช้เรียงอันดับ ไม่ใช่ความน่าจะเป็น"),
        ("ความต้องการความงามที่ AI อนุมาน", "ข้อความที่ใช้จับคู่กับสินค้า AI อนุมานจากคำค้น ไม่ใช่สิ่งที่คนค้นหาโดยตรง"),
        ("หมวดสินค้าที่ AI โยงไว้", "หมวดที่ AI โยงไว้ตอนต้นน้ำ (①-⑤) ก่อนจับคู่สินค้าจริง อาจไม่ตรงกับหมวดหลัก (L2) ที่ ⑫ เลือก"),
        ("สินค้าที่ใกล้พอ", "'ไม่มีสินค้าที่ใกล้พอ' = ใน 5 ชิ้นที่ใกล้สุด มีในหมวดหลักน้อยกว่า 2 ชิ้น"),
        ("คาดโตสะสม", "การเติบโตที่คาดของค่าเฉลี่ยการค้นหาในปีนั้น เทียบกับ 12 เดือนล่าสุด (สะสม)"),
        ("คาดโตเฉลี่ยต่อปี", "ตัวเลขชุดเดียวกับที่รายงานแสดง (+X% ต่อปีแบบทบต้นจาก 12 เดือนล่าสุดถึงปีนั้น)"),
        ("คะแนน (0-1)", "percentile ของการเติบโตที่คาดของเทรนด์เทียบคีย์เวิร์ดทั้งหมดในรอบ - 0.5 = ค่ากลางตลาด"),
        ("หมวดของ Google", "หมวดที่ Google จัดหัวข้อค้นหานั้นไว้ (ตามที่ Google จัด ไม่ได้แก้)"),
        ("ข้อควรระวัง", "ตัวเลขการเติบโตเป็นผลทำนายจากโมเดล ยังไม่เคยตรวจกับยอดขายจริง ควรอ่านเป็นทิศทางและอันดับ"),
    ]


def write_excel(path, sheets):
    """sheets: [(ชื่อชีต, DataFrame)] - หัวตารางตรึง, ความกว้างคอลัมน์ปรับตามข้อมูล, ข้อความยาวตัดบรรทัด, ลิงก์กดเปิดได้"""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wide = {"เหตุผลที่เลือกหมวดหลัก", "เหตุผล", "ความต้องการความงามที่ AI อนุมาน", "หัวข้อ", "ค่า"}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets:
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(frame.columns, start=1):
                letter = get_column_letter(i)
                longest = max([len(str(col))] + [len(str(v)) for v in frame[col].tolist()] or [0])
                ws.column_dimensions[letter].width = min(max(12, longest + 2), 70 if col in wide else 45)
                for cell in ws[letter][1:]:
                    if col in wide or col.startswith("ชนิดสินค้าที่เด่น") or col.startswith("หมวดของสินค้า"):
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if col.startswith("ลิงก์") and isinstance(cell.value, str) and cell.value.startswith("http"):
                        cell.hyperlink = cell.value
                        cell.font = Font(color="0563C1", underline="single")
    return path


def _pick_run(output_root, name):
    if name:
        return output_root / name
    env = os.environ.get("TREND_RUN_ID")
    if env:
        return output_root / env
    runs = sorted((p for p in output_root.iterdir() if (p / CONTEXT_FILE).exists()), key=lambda p: p.stat().st_mtime)
    if not runs:
        raise SystemExit(f"❌ ไม่พบรอบที่มี {CONTEXT_FILE} ใน {output_root} - สร้างรายงาน ⑯ ก่อน")
    return runs[-1]


def build(run_dir, with_links=True):
    run_dir = Path(run_dir)
    for fname in (CONTEXT_FILE, RESULT_FILE):
        if not (run_dir / fname).exists():
            raise SystemExit(f"❌ ไม่พบ {fname} ใน {run_dir.name} - รัน ⑫ และสร้างรายงาน ⑯ ของรอบนี้ก่อน")
    ctx = json.loads((run_dir / CONTEXT_FILE).read_text(encoding="utf-8"))
    result = json.loads((run_dir / RESULT_FILE).read_text(encoding="utf-8"))
    cands = pd.read_csv(run_dir / CANDIDATES_FILE) if (run_dir / CANDIDATES_FILE).exists() else None
    urls = None
    if with_links:
        from product.euromonitor_mapping import load_catalog
        cat = load_catalog(beauty_only=False)
        urls = {(b, n): u for b, n, u in zip(cat["brand"], cat["product_name"], cat["product_url"])
                if isinstance(u, str) and u.startswith("http")}
    sheets = [(SHEET_SUMMARY, pd.DataFrame(summary_rows(ctx, result))),
              (SHEET_TYPES, pd.DataFrame(type_rows(ctx, result))),
              (SHEET_PRODUCTS, pd.DataFrame(product_rows(ctx, result, urls)))]
    google = google_rows(ctx, cands)
    if google:
        sheets.append((SHEET_GOOGLE, pd.DataFrame(google)))
    sheets.append((SHEET_NOTES, pd.DataFrame(notes_rows(ctx, result), columns=["หัวข้อ", "ค่า"])))
    return write_excel(run_dir / OUTPUT_FILE, sheets), {n: len(f) for n, f in sheets}


def main():
    ap = argparse.ArgumentParser(description="Excel รายเทรนด์ x หมวดสินค้าของรอบ")
    ap.add_argument("--run", default=None, help="ชื่อโฟลเดอร์รอบใน main/output (ไม่ใส่ = TREND_RUN_ID หรือรอบล่าสุด)")
    ap.add_argument("--no-links", action="store_true", help="ไม่โหลดแคตตาล็อกเพื่อหาลิงก์สินค้า (เร็วกว่า)")
    a = ap.parse_args()
    root = Path(__file__).resolve().parents[2] / "output"
    path, sizes = build(_pick_run(root, a.run), with_links=not a.no_links)
    print(f"✅ บันทึก {path}")
    print("   " + " · ".join(f"{n} {k} แถว" for n, k in sizes.items()))


if __name__ == "__main__":
    main()
